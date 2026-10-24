import json
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Generator
from uuid import UUID, uuid4

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
import pytest
import reportlab
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.audits import get_audit_enqueue
from app.api.routes.documents import get_document_storage
from app.core.config import settings
from app.db.models import (
    Audit,
    Document,
    Feedback,
    Finding,
    FindingCitation,
    LegalProvision,
    LegalSource,
    Organization,
    OrganizationMembership,
    User,
)
from app.db.session import get_db
from app.main import app
from app.services.compliance_analysis import analyze_audit
from app.services.llm_provider import LLMProviderError
from app.storage.interface import LocalPrivateFileStorage
from app.workers.audit_worker import run_audit_job


class MockAnalyzer:
    model_name = "test-only-mock-analyzer"
    model_version = "fixture-v1"

    def __init__(self, fail: bool = False) -> None:
        self.fail = fail

    def analyze_clause(self, clause: dict, evidence: list[dict]) -> str:
        if self.fail:
            raise LLMProviderError("synthetic test-only provider failure")
        return json.dumps(
            {
                "category": "insufficient_evidence",
                "severity": "informational",
                "explanation": "The test fixture has no authoritative legal evidence.",
                "recommended_action": "Check the applicable law with qualified counsel.",
                "evidence_assessment": "insufficient",
                "citations": [],
            }
        )


class AuditTestStorage(LocalPrivateFileStorage):
    pass


def _docx_bytes(text: str = "The employee salary will be paid monthly.") -> bytes:
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _register(client: TestClient, email: str) -> tuple[dict, str]:
    response = client.post(
        "/api/v1/auth/register",
        json={
            "name": "Audit Test User",
            "email": email,
            "password": "a-long-audit-test-password",
        },
    )
    assert response.status_code == 201
    login = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "a-long-audit-test-password"},
    )
    assert login.status_code == 200
    return response.json(), login.json()["access_token"]


def _upload_docx(client: TestClient, token: str) -> dict:
    response = client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "employment-contract.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201
    return response.json()


def _create_completed_report(
    client: TestClient,
    db: Session,
    email: str,
    *,
    clause_text: str = "The employee salary will be paid monthly.",
) -> tuple[dict, str, Audit, Finding]:
    registration, token = _register(client, email)
    user_id = UUID(registration["id"])
    uploaded = _upload_docx(client, token)
    document = db.get(Document, UUID(uploaded["id"]))
    user = db.get(User, user_id)
    assert document is not None and user is not None

    source = LegalSource(
        title="Historical source title",
        law_number="fixture-only",
        version_label="fixture-version-1",
        source_type="statute",
        jurisdiction="ET",
        language="en",
        legal_status="current",
        verification_status="verified",
        verified_by=user.id,
        verified_at=datetime.now(timezone.utc),
        original_source_url="https://example.test/fixture-law.pdf",
        file_checksum="f" * 64,
        source_file_storage_key=f"legal-sources/{uuid4()}.pdf",
    )
    provision = LegalProvision(
        legal_source=source,
        article_number="Fixture Article 4",
        sub_article_reference="(a)",
        section_title="Fixture heading",
        text="Historical provision text from a test-only synthetic fixture.",
        language="en",
    )
    now = datetime.now(timezone.utc)
    audit = Audit(
        document=document,
        requester=user,
        audit_status="completed",
        review_options={"jurisdiction": "ET", "language": "en"},
        processing_warnings=["Test-only extraction warning."],
        analysis_version="test-report-v1",
        model_identifier="test-only-model:fixture-v1",
        started_at=now,
        completed_at=now,
    )
    finding = Finding(
        audit=audit,
        finding_type="potential_legal_conflict:compensation:clause-1",
        severity="medium",
        clause_text=clause_text,
        explanation="Test-only explanation of a potential concern.",
        recommended_action="Ask qualified counsel to review this test fixture.",
        evidence_status="supported",
    )
    citation = FindingCitation(
        finding=finding,
        legal_provision=provision,
        supporting_quote="Historical provision text",
        relevance_information="test-only fixture",
        source_title_snapshot="Historical source title",
        source_version_snapshot="fixture-version-1",
        source_url_snapshot="https://example.test/fixture-law.pdf",
        article_number_snapshot="Fixture Article 4",
        sub_article_reference_snapshot="(a)",
        provision_text_snapshot=(
            "Historical provision text from a test-only synthetic fixture."
        ),
    )
    db.add_all([audit, source, citation])
    db.commit()
    return registration, token, audit, finding


@pytest.fixture
def audit_client(
    db_session: Session,
    tmp_path: Path,
) -> Generator[tuple[TestClient, Session, LocalPrivateFileStorage, list[UUID]], None, None]:
    storage = AuditTestStorage(tmp_path / "private-files")
    enqueued: list[UUID] = []

    def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_document_storage] = lambda: storage
    app.dependency_overrides[get_audit_enqueue] = lambda: (
        lambda audit_id: enqueued.append(audit_id)
    )
    try:
        with TestClient(app) as client:
            yield client, db_session, storage, enqueued
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_document_storage, None)
        app.dependency_overrides.pop(get_audit_enqueue, None)


def test_uploaded_document_audit_lifecycle_and_duplicate_delivery(
    audit_client,
    monkeypatch,
) -> None:
    client, db, storage, enqueued = audit_client
    _, token = _register(client, "audit-lifecycle@example.com")
    uploaded = _upload_docx(client, token)
    created = client.post(
        "/api/v1/audits",
        headers={"Authorization": f"Bearer {token}"},
        json={
            "document_id": uploaded["id"],
            "review_options": {"review_date": "2026-10-24", "jurisdiction": "ET"},
        },
    )
    assert created.status_code == 202
    audit_id = UUID(created.json()["id"])
    assert created.json()["audit_status"] == "queued"
    assert enqueued == [audit_id]
    assert client.get(
        f"/api/v1/audits/{audit_id}/status",
        headers={"Authorization": f"Bearer {token}"},
    ).json()["status"] == "queued"
    assert client.get(
        f"/api/v1/audits/{audit_id}/findings",
        headers={"Authorization": f"Bearer {token}"},
    ).status_code == 409

    seen_states: list[str] = []
    original_read = storage.read

    def checked_read(key: str) -> bytes:
        seen_states.append(db.get(Audit, audit_id).audit_status)
        return original_read(key)

    storage.read = checked_read
    import app.services.compliance_analysis as analysis_module

    original_retrieve = analysis_module.retrieve_legal_evidence

    def checked_retrieve(*args, **kwargs):
        seen_states.append(db.get(Audit, audit_id).audit_status)
        return original_retrieve(*args, **kwargs)

    monkeypatch.setattr(analysis_module, "retrieve_legal_evidence", checked_retrieve)
    analyzer = MockAnalyzer()
    original_analyze = analyzer.analyze_clause

    def checked_analyze_clause(clause: dict, evidence: list[dict]) -> str:
        seen_states.append(db.get(Audit, audit_id).audit_status)
        return original_analyze(clause, evidence)

    analyzer.analyze_clause = checked_analyze_clause

    original_persist = analysis_module._persist_finding

    def checked_persist(*args, **kwargs):
        seen_states.append(db.get(Audit, audit_id).audit_status)
        return original_persist(*args, **kwargs)

    monkeypatch.setattr(analysis_module, "_persist_finding", checked_persist)
    assert run_audit_job(
        audit_id,
        db=db,
        storage=storage,
        llm_provider=analyzer,
    )
    assert not run_audit_job(
        audit_id,
        db=db,
        storage=storage,
        llm_provider=analyzer,
    )
    assert seen_states == [
        "extracting",
        "retrieving_sources",
        "analyzing",
        "generating_report",
    ]
    db.expire_all()
    audit = db.get(Audit, audit_id)
    assert audit.audit_status == "completed"
    assert audit.completed_at is not None
    assert audit.analysis_version == "1"
    assert audit.model_identifier == "test-only-mock-analyzer:fixture-v1"
    assert db.scalar(select(Finding).where(Finding.audit_id == audit_id)) is not None

    findings_response = client.get(
        f"/api/v1/audits/{audit_id}/findings",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert findings_response.status_code == 200
    findings = findings_response.json()["findings"]
    assert len(findings) == 1
    assert findings[0]["category"] == "insufficient_evidence"
    assert findings[0]["citations"] == []


def test_audit_endpoints_hide_another_users_audit(audit_client) -> None:
    client, _, _, _ = audit_client
    _, owner_token = _register(client, "audit-owner@example.com")
    _, stranger_token = _register(client, "audit-stranger@example.com")
    uploaded = _upload_docx(client, owner_token)
    created = client.post(
        "/api/v1/audits",
        headers={"Authorization": f"Bearer {owner_token}"},
        json={"document_id": uploaded["id"]},
    )
    audit_id = created.json()["id"]

    for path in (
        f"/api/v1/audits/{audit_id}",
        f"/api/v1/audits/{audit_id}/status",
        f"/api/v1/audits/{audit_id}/findings",
    ):
        response = client.get(
            path,
            headers={"Authorization": f"Bearer {stranger_token}"},
        )
        assert response.status_code == 404

    stranger_audits = client.get(
        "/api/v1/audits",
        headers={"Authorization": f"Bearer {stranger_token}"},
    )
    assert stranger_audits.status_code == 200
    assert stranger_audits.json()["items"] == []


def test_audit_creation_rejects_other_users_document(audit_client) -> None:
    client, _, _, _ = audit_client
    _, owner_token = _register(client, "audit-doc-owner@example.com")
    _, stranger_token = _register(client, "audit-doc-stranger@example.com")
    uploaded = _upload_docx(client, owner_token)
    response = client.post(
        "/api/v1/audits",
        headers={"Authorization": f"Bearer {stranger_token}"},
        json={"document_id": uploaded["id"]},
    )
    assert response.status_code == 404


def test_organization_audits_require_membership(audit_client) -> None:
    client, db, _, _ = audit_client
    _, owner_token = _register(client, "audit-org-owner@example.com")
    _, member_token = _register(client, "audit-org-member@example.com")
    _, outsider_token = _register(client, "audit-org-outsider@example.com")
    owner_response = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {owner_token}"},
    )
    member_response = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {member_token}"},
    )
    owner = db.get(User, UUID(owner_response.json()["id"]))
    member = db.get(User, UUID(member_response.json()["id"]))
    organization = Organization(name="Audit Org", organization_type="employer")
    db.add(organization)
    db.flush()
    db.add(
        OrganizationMembership(
            organization_id=organization.id,
            user_id=member.id,
            membership_role="member",
        )
    )
    outsider = db.scalar(select(User).where(User.email == "audit-org-outsider@example.com"))
    assert outsider is not None
    db.add(
        OrganizationMembership(
            organization_id=organization.id,
            user_id=outsider.id,
            membership_role="viewer",
        )
    )
    document = Document(
        owner=owner,
        organization=organization,
        original_filename="org-contract.docx",
        storage_key="docs/org-contract",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=10,
        document_type="docx",
    )
    audit = Audit(document=document, requester=member)
    db.add(audit)
    db.commit()

    accessible = client.get(
        f"/api/v1/audits/{audit.id}",
        headers={"Authorization": f"Bearer {member_token}"},
    )
    denied = client.get(
        f"/api/v1/audits/{audit.id}/findings",
        headers={"Authorization": f"Bearer {outsider_token}"},
    )
    assert accessible.status_code == 200
    assert denied.status_code == 403


def test_unauthenticated_users_cannot_access_audit_or_findings(audit_client) -> None:
    client, db, _, _ = audit_client
    owner = User(name="Protected Audit User", email="protected-audit@example.com", password_hash="test")
    db.add(owner)
    db.flush()
    document = Document(
        owner=owner,
        original_filename="protected.docx",
        storage_key="docs/protected",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=10,
        document_type="docx",
    )
    audit = Audit(document=document, requester=owner)
    db.add(audit)
    db.commit()

    for path in (
        "/api/v1/audits",
        f"/api/v1/audits/{audit.id}",
        f"/api/v1/audits/{audit.id}/status",
        f"/api/v1/audits/{audit.id}/findings",
    ):
        assert client.get(path).status_code == 401

    assert client.post(
        "/api/v1/audits",
        json={"document_id": str(document.id)},
    ).status_code == 401


def test_audit_pagination_bounds_are_validated(audit_client) -> None:
    client, _, _, _ = audit_client
    _, token = _register(client, "audit-pagination@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    assert client.get(
        "/api/v1/audits?limit=101",
        headers=headers,
    ).status_code == 422
    assert client.get(
        "/api/v1/audits?offset=1000001",
        headers=headers,
    ).status_code == 422


def test_report_uses_saved_results_and_citation_snapshots_privately(
    audit_client,
) -> None:
    client, db, _, _ = audit_client
    registration, token, audit, finding = _create_completed_report(
        client,
        db,
        "audit-report-owner@example.com",
    )
    citation = db.query(FindingCitation).filter_by(finding_id=finding.id).one()
    source = db.get(LegalSource, citation.legal_provision.legal_source_id)
    provision = citation.legal_provision
    assert source is not None
    source.title = "Changed later source title"
    source.version_label = "later-version"
    source.original_source_url = "https://example.test/later.pdf"
    provision.article_number = "Later article"
    provision.text = "Changed provision content"
    db.commit()

    report_response = client.get(
        f"/api/v1/audits/{audit.id}/report",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert report_response.status_code == 200
    assert report_response.headers["cache-control"] == "private, no-store"
    report = report_response.json()
    assert report["audit_id"] == str(audit.id)
    assert report["audit_created_at"]
    assert report["document"]["original_filename"] == "employment-contract.docx"
    assert report["review_scope"]["jurisdiction"] == "ET"
    assert report["summary"]["finding_count"] == 1
    assert report["summary"]["severity_counts"] == {"medium": 1}
    assert report["findings"][0]["clause_excerpt"] == (
        "The employee salary will be paid monthly."
    )
    assert report["findings"][0]["explanation"].startswith("Test-only")
    assert report["findings"][0]["citations"][0]["source_title"] == (
        "Historical source title"
    )
    assert report["findings"][0]["citations"][0]["source_version"] == (
        "fixture-version-1"
    )
    assert report["findings"][0]["citations"][0]["article_number"] == (
        "Fixture Article 4"
    )
    assert report["findings"][0]["citations"][0]["provision_text"].startswith(
        "Historical provision text"
    )
    assert report["processing_warnings"] == ["Test-only extraction warning."]
    assert any(
        "AI-assisted review" in item
        for item in report["limitations"]["items"]
    )
    assert "qualified Ethiopian legal professional" in (
        report["limitations"]["disclaimer"]
    )

    pdf_response = client.get(
        f"/api/v1/audits/{audit.id}/report.pdf",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert pdf_response.status_code == 200
    assert pdf_response.headers["content-type"] == "application/pdf"
    assert pdf_response.headers["cache-control"] == "private, no-store"
    assert pdf_response.content.startswith(b"%PDF-")
    assert f"legalcheck-audit-{audit.id}.pdf" in pdf_response.headers[
        "content-disposition"
    ]

    _, stranger_token = _register(client, "audit-report-stranger@example.com")
    unauthorized_report = client.get(
        f"/api/v1/audits/{audit.id}/report",
        headers={"Authorization": f"Bearer {stranger_token}"},
    )
    unauthorized_pdf = client.get(
        f"/api/v1/audits/{audit.id}/report.pdf",
        headers={"Authorization": f"Bearer {stranger_token}"},
    )
    assert unauthorized_report.status_code == 404
    assert unauthorized_pdf.status_code == 404
    assert client.get(f"/api/v1/audits/{audit.id}/report").status_code == 401
    assert client.get(f"/api/v1/audits/{audit.id}/report.pdf").status_code == 401


def test_report_preserves_unicode_and_requires_font_for_ethiopic(audit_client) -> None:
    client, db, _, _ = audit_client
    _, token, audit, _ = _create_completed_report(
        client,
        db,
        "audit-report-unicode@example.com",
        clause_text="የሥራ ውል።",
    )

    response = client.get(
        f"/api/v1/audits/{audit.id}/report",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 200
    assert response.json()["findings"][0]["clause_excerpt"] == "የሥራ ውል።"

    pdf_response = client.get(
        f"/api/v1/audits/{audit.id}/report.pdf",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert pdf_response.status_code == 503
    assert "REPORT_UNICODE_FONT_PATH" in pdf_response.json()["error"]["message"]


def test_unicode_font_configuration_generates_report_with_embedded_true_type_font(
    audit_client,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    client, db, _, _ = audit_client
    _, token, audit, _ = _create_completed_report(
        client,
        db,
        "audit-report-accented@example.com",
        clause_text="The employee signed at the café.",
    )
    vera_font = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    monkeypatch.setattr(settings, "report_unicode_font_path", vera_font)

    response = client.get(
        f"/api/v1/audits/{audit.id}/report.pdf",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
    assert b"/FontFile2" in response.content


def test_feedback_validates_category_and_description_and_checks_ownership(
    audit_client,
) -> None:
    client, db, _, _ = audit_client
    _, token, audit, finding = _create_completed_report(
        client,
        db,
        "audit-feedback-owner@example.com",
    )
    feedback_url = f"/api/v1/findings/{finding.id}/feedback"
    owner_headers = {"Authorization": f"Bearer {token}"}
    valid = client.post(
        feedback_url,
        headers=owner_headers,
        json={
            "feedback_type": "incorrect_citation",
            "description": "Please verify the cited source version.",
        },
    )
    assert valid.status_code == 201
    assert valid.json()["finding_id"] == str(finding.id)
    assert valid.json()["feedback_type"] == "incorrect_citation"
    assert valid.json()["description"] == "Please verify the cited source version."

    invalid_category = client.post(
        feedback_url,
        headers=owner_headers,
        json={"feedback_type": "made_up"},
    )
    missing_other_description = client.post(
        feedback_url,
        headers=owner_headers,
        json={"feedback_type": "other"},
    )
    too_long = client.post(
        feedback_url,
        headers=owner_headers,
        json={
            "feedback_type": "missing_issue",
            "description": "x" * 5001,
        },
    )
    assert invalid_category.status_code == 422
    assert missing_other_description.status_code == 422
    assert too_long.status_code == 422

    _, stranger_token = _register(client, "audit-feedback-stranger@example.com")
    unauthorized = client.post(
        feedback_url,
        headers={"Authorization": f"Bearer {stranger_token}"},
        json={"feedback_type": "incorrect_finding"},
    )
    assert unauthorized.status_code == 404
    assert db.query(Feedback).count() == 1
    assert client.post(
        feedback_url,
        json={"feedback_type": "incorrect_finding"},
    ).status_code == 401


def test_report_and_feedback_require_completed_audit(audit_client) -> None:
    client, db, _, _ = audit_client
    registration, token = _register(client, "audit-report-pending@example.com")
    uploaded = _upload_docx(client, token)
    document = db.get(Document, UUID(uploaded["id"]))
    owner = db.get(User, UUID(registration["id"]))
    assert document is not None and owner is not None
    audit = Audit(document=document, requester=owner, audit_status="queued")
    finding = Finding(
        audit=audit,
        finding_type="informational_observation:test:clause-1",
        severity="informational",
        clause_text="Test clause",
        explanation="Pending",
    )
    db.add(finding)
    db.commit()

    headers = {"Authorization": f"Bearer {token}"}
    assert client.get(
        f"/api/v1/audits/{audit.id}/report",
        headers=headers,
    ).status_code == 409
    assert client.get(
        f"/api/v1/audits/{audit.id}/report.pdf",
        headers=headers,
    ).status_code == 409
    assert client.post(
        f"/api/v1/findings/{finding.id}/feedback",
        headers=headers,
        json={"feedback_type": "incorrect_finding"},
    ).status_code == 409


def test_invalid_transition_and_processing_failure_do_not_complete(audit_client) -> None:
    _, db, storage, _ = audit_client
    owner = User(name="Transition User", email="transition@example.com", password_hash="test")
    db.add(owner)
    db.flush()
    document = Document(
        owner=owner,
        original_filename="contract.docx",
        storage_key="docs/transition",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=len(_docx_bytes()),
        document_type="docx",
    )
    audit = Audit(document=document, requester=owner, audit_status="queued")
    db.add(audit)
    db.commit()
    try:
        analyze_audit(
            audit.id,
            db=db,
            storage=storage,
            llm_provider=MockAnalyzer(),
            initial_status="analyzing",
        )
    except Exception as exc:
        assert "processable state" in str(exc)
    else:
        raise AssertionError("The invalid initial state should have been rejected.")
    assert db.get(Audit, audit.id).audit_status == "queued"

    storage.store("docs/transition", _docx_bytes())
    assert run_audit_job(
        audit.id,
        db=db,
        storage=storage,
        llm_provider=MockAnalyzer(fail=True),
    )
    db.refresh(audit)
    assert audit.audit_status == "failed"
    assert audit.completed_at is None
    assert audit.failure_code == "llm_provider_failure"
    assert db.scalar(select(Finding).where(Finding.audit_id == audit.id)) is None


def test_old_worker_attempt_cannot_overwrite_retried_attempt(audit_client) -> None:
    _, db, storage, _ = audit_client
    owner = User(name="Attempt User", email="attempt@example.com", password_hash="test")
    db.add(owner)
    db.flush()
    document = Document(
        owner=owner,
        original_filename="contract.docx",
        storage_key="docs/attempt",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_size=len(_docx_bytes()),
        document_type="docx",
    )
    old_attempt = uuid4()
    audit = Audit(
        document=document,
        requester=owner,
        audit_status="extracting",
        processing_attempt_id=old_attempt,
    )
    db.add(audit)
    db.commit()
    new_attempt = uuid4()

    class SupersededAttemptStorage:
        def read(self, key: str) -> bytes:
            current = db.get(Audit, audit.id)
            current.processing_attempt_id = new_attempt
            db.commit()
            return _docx_bytes()

        def store(self, key: str, content: bytes) -> None:
            raise NotImplementedError

        def delete(self, key: str) -> None:
            raise NotImplementedError

    try:
        analyze_audit(
            audit.id,
            db=db,
            storage=SupersededAttemptStorage(),
            llm_provider=MockAnalyzer(),
            initial_status="extracting",
            attempt_id=old_attempt,
        )
    except Exception as exc:
        assert "could not be completed" in str(exc)
    else:
        raise AssertionError("The stale worker attempt should have stopped.")
    db.refresh(audit)
    assert audit.audit_status == "extracting"
    assert audit.processing_attempt_id == new_attempt


def test_failed_audit_can_be_retried_and_queued_audit_cannot(audit_client) -> None:
    client, db, _, enqueued = audit_client
    owner, token = _register(client, "audit-retry@example.com")
    uploaded = _upload_docx(client, token)
    document = db.get(Document, UUID(uploaded["id"]))
    failed = Audit(
        document=document,
        requested_by=UUID(owner["id"]),
        audit_status="failed",
        failure_code="analysis_failed",
        failure_message="Safe failure.",
    )
    queued = Audit(
        document=document,
        requested_by=UUID(owner["id"]),
        audit_status="queued",
    )
    db.add_all([failed, queued])
    db.commit()

    response = client.post(
        f"/api/v1/audits/{failed.id}/retry",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 202
    assert response.json()["audit_status"] == "queued"
    assert response.json()["failure_code"] is None
    assert UUID(response.json()["id"]) in enqueued

    rejected = client.post(
        f"/api/v1/audits/{queued.id}/retry",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert rejected.status_code == 409


def test_queue_failure_is_recorded_and_does_not_leave_a_queued_audit(audit_client) -> None:
    client, db, _, _ = audit_client
    _, token = _register(client, "audit-queue-failure@example.com")
    uploaded = _upload_docx(client, token)

    def fail_enqueue(_audit_id: UUID) -> None:
        raise ConnectionError("synthetic test-only Redis outage")

    app.dependency_overrides[get_audit_enqueue] = lambda: fail_enqueue
    response = client.post(
        "/api/v1/audits",
        headers={"Authorization": f"Bearer {token}"},
        json={"document_id": uploaded["id"]},
    )

    assert response.status_code == 503
    audit = db.scalar(select(Audit).where(Audit.document_id == UUID(uploaded["id"])))
    assert audit is not None
    assert audit.audit_status == "failed"
    assert audit.failure_code == "queue_unavailable"


def test_stale_worker_state_can_be_recovered(audit_client) -> None:
    client, db, _, enqueued = audit_client
    owner, token = _register(client, "audit-stale@example.com")
    uploaded = _upload_docx(client, token)
    audit = Audit(
        document_id=UUID(uploaded["id"]),
        requested_by=UUID(owner["id"]),
        audit_status="analyzing",
        started_at=datetime.now(timezone.utc)
        - timedelta(seconds=settings.audit_stale_after_seconds + 60),
    )
    db.add(audit)
    db.commit()

    response = client.post(
        f"/api/v1/audits/{audit.id}/retry",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 202
    assert response.json()["audit_status"] == "queued"
    assert "stale worker" in response.json()["processing_warnings"][-1].lower()
    assert UUID(response.json()["id"]) in enqueued
