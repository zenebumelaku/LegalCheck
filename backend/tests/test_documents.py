from io import BytesIO
from pathlib import Path
from typing import Generator
from uuid import UUID

import fitz
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.documents import get_document_storage
from app.core.config import settings
from app.db.models import (
    Audit,
    Document,
    Organization,
    OrganizationMembership,
    User,
)
from app.db.session import get_db
from app.main import app
from app.storage.interface import LocalPrivateFileStorage


@pytest.fixture
def document_client(
    db_session: Session,
    tmp_path: Path,
) -> Generator[TestClient, None, None]:
    def override_get_db():
        yield db_session

    storage = LocalPrivateFileStorage(tmp_path / "private-files")
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_document_storage] = lambda: storage
    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_document_storage, None)


def _register_and_token(client: TestClient, email: str) -> tuple[dict, str]:
    registration = client.post(
        "/api/v1/auth/register",
        json={
            "name": "Document User",
            "email": email,
            "password": "a-long-upload-test-password",
        },
    )
    assert registration.status_code == 201
    login = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "a-long-upload-test-password"},
    )
    assert login.status_code == 200
    return registration.json(), login.json()["access_token"]


def _pdf_bytes(text: str = "Employment contract paragraph for extraction.") -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def _docx_bytes(text: str = "DOCX employment paragraph.") -> bytes:
    document = DocxDocument()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Contract table clause."
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def _upload(
    client: TestClient,
    token: str,
    filename: str,
    content: bytes,
    content_type: str,
):
    return client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": (filename, content, content_type)},
    )


def test_upload_pdf_extracts_text_and_keeps_storage_private(
    document_client: TestClient,
    tmp_path: Path,
) -> None:
    _, token = _register_and_token(document_client, "pdf@example.com")
    pdf_content = _pdf_bytes()

    response = _upload(
        document_client,
        token,
        "../../Employment contract.pdf",
        pdf_content,
        "application/pdf",
    )

    assert response.status_code == 201
    body = response.json()
    assert body["document_type"] == "pdf"
    assert body["original_filename"] == "Employment contract.pdf"
    assert body["pages"][0]["page_number"] == 1
    assert body["pages"][0]["paragraphs"][0]["paragraph_number"] >= 1
    assert "Employment contract paragraph" in body["pages"][0]["paragraphs"][0]["text"]
    assert body["warnings"] == []
    assert "storage_key" not in body
    assert "private-files" not in response.text
    stored_files = list((tmp_path / "private-files").rglob("*.pdf"))
    assert len(stored_files) == 1
    assert stored_files[0].name != "Employment contract.pdf"
    assert stored_files[0].read_bytes() == pdf_content


def test_upload_docx_preserves_paragraph_references_without_fake_page_number(
    document_client: TestClient,
) -> None:
    _, token = _register_and_token(document_client, "docx@example.com")

    response = _upload(
        document_client,
        token,
        "contract.docx",
        _docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 201
    page = response.json()["pages"][0]
    assert page["page_number"] is None
    assert page["paragraphs"][0]["paragraph_number"] == 1
    assert page["paragraphs"][0]["text"] == "DOCX employment paragraph."
    assert page["paragraphs"][1]["text"] == "Contract table clause."


def test_scan_like_pdf_returns_ocr_warning(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "scan@example.com")

    response = _upload(
        document_client,
        token,
        "scan.pdf",
        _pdf_bytes(text=""),
        "application/pdf",
    )

    assert response.status_code == 201
    assert response.json()["processing_status"] == "warning"
    assert any("scanned" in warning.lower() for warning in response.json()["warnings"])


def test_upload_rejects_unsupported_extension(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "extension@example.com")

    response = _upload(document_client, token, "contract.txt", b"text", "text/plain")

    assert response.status_code == 415


def test_document_routes_require_authentication(document_client: TestClient) -> None:
    response = document_client.get("/api/v1/documents")

    assert response.status_code == 401


def test_upload_rejects_mismatched_mime_and_signature(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "mismatch@example.com")

    wrong_mime = _upload(
        document_client,
        token,
        "contract.pdf",
        _pdf_bytes(),
        "text/plain",
    )
    wrong_signature = _upload(
        document_client,
        token,
        "contract.pdf",
        b"PK\x03\x04not a PDF",
        "application/pdf",
    )

    assert wrong_mime.status_code == 415
    assert wrong_signature.status_code == 415


def test_upload_rejects_empty_and_corrupted_files(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "corrupt@example.com")

    empty = _upload(document_client, token, "empty.pdf", b"", "application/pdf")
    corrupted = _upload(
        document_client,
        token,
        "corrupted.pdf",
        b"%PDF-1.7\nnot a valid PDF body",
        "application/pdf",
    )

    assert empty.status_code == 400
    assert corrupted.status_code == 422


def test_upload_rejects_corrupted_and_encrypted_docx(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "docx-invalid@example.com")
    content_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    corrupted = _upload(
        document_client,
        token,
        "broken.docx",
        b"PK\x03\x04not a valid zip archive",
        content_type,
    )
    encrypted = _upload(
        document_client,
        token,
        "encrypted.docx",
        bytes.fromhex("D0CF11E0A1B11AE1") + (b"\x00" * 32),
        content_type,
    )

    assert corrupted.status_code == 422
    assert encrypted.status_code == 422


def test_upload_rejects_encrypted_pdf(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "encrypted@example.com")
    document = fitz.open()
    document.new_page()
    encrypted = document.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-password",
        user_pw="user-password",
    )
    document.close()

    response = _upload(
        document_client,
        token,
        "encrypted.pdf",
        encrypted,
        "application/pdf",
    )

    assert response.status_code == 422
    assert "Password-protected" in response.json()["error"]["message"]


def test_upload_rejects_configured_oversize_file(
    document_client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, token = _register_and_token(document_client, "large@example.com")
    monkeypatch.setattr(settings, "max_upload_size_mb", 1)
    content = b"%PDF-" + (b"x" * (1024 * 1024))

    response = _upload(
        document_client,
        token,
        "large.pdf",
        content,
        "application/pdf",
    )

    assert response.status_code == 413


def test_upload_rejects_oversized_request_body(document_client: TestClient) -> None:
    _, token = _register_and_token(document_client, "large-request@example.com")
    content = b"%PDF-" + (b"x" * (11 * 1024 * 1024))

    response = _upload(
        document_client,
        token,
        "large-request.pdf",
        content,
        "application/pdf",
    )

    assert response.status_code == 413


def test_cross_user_cannot_read_list_or_delete_document(
    document_client: TestClient,
    db_session: Session,
) -> None:
    first_user, first_token = _register_and_token(document_client, "owner@example.com")
    _, other_token = _register_and_token(document_client, "other@example.com")
    uploaded = _upload(
        document_client,
        first_token,
        "owned.docx",
        _docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    document_id = uploaded.json()["id"]
    other_headers = {"Authorization": f"Bearer {other_token}"}

    detail = document_client.get(
        f"/api/v1/documents/{document_id}",
        headers=other_headers,
    )
    deleted = document_client.delete(
        f"/api/v1/documents/{document_id}",
        headers=other_headers,
    )
    listing = document_client.get("/api/v1/documents", headers=other_headers)

    assert detail.status_code == 404
    assert deleted.status_code == 404
    assert listing.json()["items"] == []
    stored = db_session.scalar(select(Document).where(Document.id == UUID(document_id)))
    assert stored is not None
    assert str(stored.owner_user_id) == first_user["id"]


def test_delete_soft_deletes_metadata_and_removes_file_without_losing_audit(
    document_client: TestClient,
    db_session: Session,
    tmp_path: Path,
) -> None:
    user_data, token = _register_and_token(document_client, "delete@example.com")
    uploaded = _upload(
        document_client,
        token,
        "delete.docx",
        _docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    document_id = UUID(uploaded.json()["id"])
    document = db_session.get(Document, document_id)
    user = db_session.get(User, UUID(user_data["id"]))
    assert document is not None and user is not None
    audit = Audit(document=document, requester=user)
    db_session.add(audit)
    db_session.commit()
    storage_files = list((tmp_path / "private-files").rglob("*.docx"))
    assert len(storage_files) == 1

    response = document_client.delete(
        f"/api/v1/documents/{document_id}",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 204
    assert document.deleted_at is not None
    assert document.processing_status == "deleted"
    assert db_session.get(Audit, audit.id) is not None
    assert not storage_files[0].exists()
    assert document_client.get(
        f"/api/v1/documents/{document_id}",
        headers={"Authorization": f"Bearer {token}"},
    ).status_code == 404


def test_organization_documents_require_membership_for_access(
    document_client: TestClient,
    db_session: Session,
) -> None:
    owner_data, owner_token = _register_and_token(document_client, "org-owner@example.com")
    visitor_data, visitor_token = _register_and_token(
        document_client,
        "org-visitor@example.com",
    )
    owner = db_session.get(User, UUID(owner_data["id"]))
    assert owner is not None
    organization = Organization(name="Shared Org", organization_type="employer")
    membership = OrganizationMembership(
        organization=organization,
        user=owner,
        membership_role="member",
    )
    db_session.add(membership)
    db_session.commit()

    upload = document_client.post(
        f"/api/v1/documents/upload?organization_id={organization.id}",
        headers={"Authorization": f"Bearer {owner_token}"},
        files={
            "file": (
                "org.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload.json()["id"]
    visitor_headers = {"Authorization": f"Bearer {visitor_token}"}

    denied = document_client.get(
        f"/api/v1/documents/{document_id}",
        headers=visitor_headers,
    )
    assert denied.status_code == 404

    visitor = db_session.get(User, UUID(visitor_data["id"]))
    assert visitor is not None
    db_session.add(
        OrganizationMembership(
            organization=organization,
            user=visitor,
            membership_role="member",
        )
    )
    db_session.commit()

    allowed = document_client.get(
        f"/api/v1/documents/{document_id}",
        headers=visitor_headers,
    )
    assert allowed.status_code == 200
    assert allowed.json()["organization_id"] == str(organization.id)
    assert document_client.delete(
        f"/api/v1/documents/{document_id}",
        headers=visitor_headers,
    ).status_code == 403
